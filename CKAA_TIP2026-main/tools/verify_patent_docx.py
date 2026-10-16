"""Verify the generated patent .docx: equations, images, tables, and key figures."""

from __future__ import annotations

import os.path as osp
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs_patent" / "航空燃油齿轮泵故障检测算法专利.docx"

MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

EXPECTED_FIGURES = [
    "95.99", "97.92", "2.37", "73.86", "72.82", "69.41", "46.26",
    "98.6", "94.0", "62.6", "88.89", "45.54", "23.87", "2.35%",
]


def main() -> None:
    document = Document(str(DOCX))

    body = document.element.body
    omath_count = len(body.findall(f".//{MATH_NS}oMath"))
    omath_para = len(body.findall(f".//{MATH_NS}oMathPara"))
    drawings = len(body.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    ))

    text = "\n".join(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)

    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]

    print(f"file            : {DOCX.name}  ({DOCX.stat().st_size:,} bytes)")
    print(f"paragraphs      : {len(document.paragraphs)}")
    print(f"tables          : {len(document.tables)}")
    print(f"images embedded : {drawings}")
    print(f"OMML equations  : {omath_count}  (in {omath_para} equation paragraphs)")
    print(f"headings        : {len(headings)}")
    print()
    print("headings:")
    for heading in headings:
        print(f"  {heading}")
    print()
    missing = [value for value in EXPECTED_FIGURES if value not in text]
    if missing:
        print(f"MISSING KEY FIGURES: {missing}")
    else:
        print(f"all {len(EXPECTED_FIGURES)} key figures present in body text/tables")

    claims = len(re.findall(r"^\s*(\d+)\.\s", text, re.MULTILINE))
    print(f"numbered claim-like lines: {claims}")


if __name__ == "__main__":
    main()
